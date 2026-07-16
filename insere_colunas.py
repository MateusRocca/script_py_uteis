"""
Código para ler uma planilha existente, identifica as colunas a partir de uma posição específica e, 
da direita para a esquerda, insere uma nova coluna em branco imediatamente após cada coluna original, 
copiando o nome do cabeçalho com o sufixo desejado, no caso em questão foi: " ajustado".
"""


from openpyxl import load_workbook

arquivo = "caminho_arquivo"

wb = load_workbook(arquivo)
ws = wb.active

linha_cabecalho = 4 # em qual linha começa o cabeçalho 
coluna_inicial = 6  # em qual coluna começa a inserir 

ultima_coluna = ws.max_column

print("Inserindo")

for col in range(ultima_coluna, coluna_inicial - 1, -1):

    nome_coluna = ws.cell(row=linha_cabecalho, column=col).value

    ws.insert_cols(col + 1)

    ws.cell(row=linha_cabecalho, column=col + 1).value = f"{nome_coluna} ajustado"

print("Salvando")

wb.save("Arquivo_ajustado.xlsx")

print("Arquivo salvo como: Arquivo_ajustado.xlsx")
